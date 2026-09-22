#!/usr/bin/env python3
"""Generate the deterministic five-city catastro demo seed from the national XLSX."""

from __future__ import annotations

import argparse
import hashlib
import re
import unicodedata
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SOURCE = Path("temp/Consolidado-Nacional-2026-publico-8 (1).xlsx")
DEFAULT_OUTPUT = Path("database/migrations/20260920_z_seed_catastro_demo.sql")


@dataclass(frozen=True)
class City:
    name: str
    province_code: str
    province_key: str
    canton_code: str
    canton_key: str
    latitude: str
    longitude: str
    urban_parishes: frozenset[str]


CITIES = (
    City(
        "Guaranda", "02", "BOLIVAR", "01", "GUARANDA", "-1.592630", "-79.000980",
        frozenset(("ANGEL POLIBIO CHAVES", "GABRIEL IGNACIO VEINTIMILLA")),
    ),
    City(
        "Riobamba", "06", "CHIMBORAZO", "01", "RIOBAMBA", "-1.663550", "-78.654650",
        frozenset(("LIZARZABURU", "MALDONADO", "VELASCO", "VELOZ")),
    ),
    City(
        "Ambato", "18", "TUNGURAHUA", "01", "AMBATO", "-1.249080", "-78.616750",
        frozenset((
            "ATOCHA FICOA", "CELIANO MONGE", "HUACHI CHICO", "HUACHI LORETO", "IZAMBA",
            "LA MERCED", "LA PENINSULA", "MATRIZ", "PISHILATA", "SAN FRANCISCO",
        )),
    ),
    City(
        "Latacunga", "05", "COTOPAXI", "01", "LATACUNGA", "-0.932650", "-78.614750",
        frozenset((
            "ELOY ALFARO SAN FELIPE", "IGNACIO FLORES PARQUE FLORES",
            "JUAN MONTALVO SAN SEBASTIAN", "LA MATRIZ", "SAN BUENAVENTURA",
        )),
    ),
    City(
        "Babahoyo", "12", "LOS RIOS", "01", "BABAHOYO", "-1.801460", "-79.534510",
        frozenset(("BARREIRO", "CLEMENTE BAQUERIZO", "DOCTOR CAMILO PONCE", "EL SALTO")),
    ),
)


def canonical(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").upper())
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^A-Z0-9]+", " ", text).strip()


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def sql_literal(value: object) -> str:
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def load_rows(source: Path) -> list[dict[str, str]]:
    worksheet = load_workbook(source, read_only=True, data_only=True).active
    iterator = worksheet.iter_rows(values_only=True)
    headers = tuple(next(iterator))
    required = (
        "RUC", "Nombre Comercial", "Número de Registro", "Actividad / Modalidad",
        "Clasificación", "Categoría", "Razón social (Propietario)", "Provincia",
        "Cantón", "Parroquia", "Estado Registro del Establecimiento",
    )
    if headers != required:
        raise ValueError(f"Unexpected XLSX headers: {headers!r}")
    index = {header: position for position, header in enumerate(headers)}
    selected: dict[str, list[dict[str, str]]] = {city.name: [] for city in CITIES}

    for raw in iterator:
        values = {header: text(raw[index[header]]) for header in headers}
        if canonical(values["Estado Registro del Establecimiento"]) != "RATIFICADO":
            continue
        if not values["Nombre Comercial"] or not values["Número de Registro"]:
            continue
        province = canonical(values["Provincia"])
        canton = canonical(values["Cantón"])
        parish = canonical(values["Parroquia"])
        for city in CITIES:
            if province == city.province_key and canton == city.canton_key and parish in city.urban_parishes:
                selected[city.name].append(values)
                break

    result: list[dict[str, str]] = []
    for city in CITIES:
        rows = sorted(selected[city.name], key=lambda row: row["Número de Registro"])
        deduplicated: dict[str, dict[str, str]] = {}
        for row in rows:
            deduplicated.setdefault(row["Número de Registro"], row)
        rows = list(deduplicated.values())[:10]
        if len(rows) != 10:
            raise ValueError(f"Expected 10 ratified urban rows for {city.name}, got {len(rows)}")
        for row in rows:
            registration = row["Número de Registro"]
            if len(registration) > 40:
                raise ValueError(f"Registration too long: {registration}")
            if row["RUC"] and (len(row["RUC"]) > 13 or not re.fullmatch(r"\d{13}", row["RUC"])):
                raise ValueError(f"Invalid RUC for selected row {registration}: {row['RUC']!r}")
            if len(row["Nombre Comercial"]) > 180:
                raise ValueError(f"Commercial name too long: {registration}")
            if len(row["Razón social (Propietario)"]) > 200:
                raise ValueError(f"Owner name too long: {registration}")
            if len(row["Actividad / Modalidad"]) > 180:
                raise ValueError(f"Activity too long: {registration}")
            if len(row["Clasificación"]) > 120 or len(row["Categoría"]) > 120:
                raise ValueError(f"Classification/category too long: {registration}")
            result.append({"city": city.name, **row})
    if len({row["Número de Registro"] for row in result}) != 50:
        raise ValueError("The selected registration numbers are not unique")
    return result


def values_sql(rows: list[dict[str, str]]) -> str:
    rendered = []
    for row in rows:
        ruc = row["RUC"] or None
        rendered.append(
            "  ("
            + ", ".join(
                sql_literal(value)
                for value in (
                    next(city.province_code for city in CITIES if city.name == row["city"]),
                    next(city.canton_code for city in CITIES if city.name == row["city"]),
                    row["city"],
                    row["Número de Registro"],
                    ruc,
                    row["Nombre Comercial"],
                    row["Razón social (Propietario)"] or None,
                    row["Actividad / Modalidad"],
                    row["Clasificación"] or None,
                    row["Categoría"] or None,
                )
            )
            + ")"
        )
    return ",\n".join(rendered)


def build_sql(source: Path, rows: list[dict[str, str]]) -> str:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    locality_values = ",\n".join(
        "  ("
        + ", ".join(
            sql_literal(value)
            for value in (
                city.province_code,
                city.canton_code,
                city.name,
                "CIUDAD",
                city.latitude,
                city.longitude,
            )
        )
        + ")"
        for city in CITIES
    )
    return f"""-- Catastro demostrativo derivado del consolidado nacional institucional.
-- Fuente: {source.as_posix()}
-- SHA-256: {digest}
-- Muestra operativa: cinco cabeceras, diez registros RATIFICADO por cabecera.
-- Las coordenadas de localidad son aproximadas; la fuente no trae coordenadas por establecimiento.

BEGIN;

SET LOCAL lock_timeout = '5s';
SET LOCAL statement_timeout = '120s';

CREATE TEMP TABLE _catastro_demo_seed (
  provincia_codigo CHAR(2) NOT NULL,
  canton_codigo CHAR(2) NOT NULL,
  localidad_nombre VARCHAR(150) NOT NULL,
  numero_registro VARCHAR(40) PRIMARY KEY,
  ruc VARCHAR(13),
  nombre_comercial VARCHAR(180) NOT NULL,
  razon_social VARCHAR(200),
  actividad VARCHAR(180) NOT NULL,
  clasificacion VARCHAR(120),
  categoria VARCHAR(120)
) ON COMMIT DROP;

INSERT INTO _catastro_demo_seed (
  provincia_codigo, canton_codigo, localidad_nombre, numero_registro, ruc,
  nombre_comercial, razon_social, actividad, clasificacion, categoria
)
VALUES
{values_sql(rows)};

INSERT INTO localidades (canton_id, nombre, tipo_localidad, latitud, longitud, ubicacion, activo)
SELECT c.id, seed.localidad_nombre, seed.tipo_localidad, seed.latitud, seed.longitud,
       ST_SetSRID(ST_MakePoint(seed.longitud, seed.latitud), 4326)::GEOGRAPHY, TRUE
FROM (
  SELECT source.provincia_codigo, source.canton_codigo, source.localidad_nombre,
         source.tipo_localidad, source.latitud::NUMERIC, source.longitud::NUMERIC
  FROM (VALUES
{locality_values}
  ) AS source(provincia_codigo, canton_codigo, localidad_nombre, tipo_localidad, latitud, longitud)
) AS seed
JOIN provincias p ON p.codigo_dpa = seed.provincia_codigo AND p.activo
JOIN cantones c ON c.provincia_id = p.id AND c.codigo_cton = seed.canton_codigo AND c.activo
ON CONFLICT (canton_id, nombre) DO UPDATE SET
  tipo_localidad = EXCLUDED.tipo_localidad,
  latitud = EXCLUDED.latitud,
  longitud = EXCLUDED.longitud,
  ubicacion = EXCLUDED.ubicacion,
  activo = TRUE;

INSERT INTO establecimientos_turisticos (
  localidad_id, numero_registro, ruc, nombre_comercial, razon_social, actividad,
  clasificacion, categoria, direccion, telefono, latitud, longitud, activo
)
SELECT l.id, seed.numero_registro, seed.ruc, seed.nombre_comercial, seed.razon_social,
       seed.actividad, seed.clasificacion, seed.categoria, NULL, NULL, NULL, NULL, TRUE
FROM _catastro_demo_seed seed
JOIN provincias p ON p.codigo_dpa = seed.provincia_codigo AND p.activo
JOIN cantones c ON c.provincia_id = p.id AND c.codigo_cton = seed.canton_codigo AND c.activo
JOIN localidades l ON l.canton_id = c.id
                    AND l.nombre = seed.localidad_nombre
                    AND l.tipo_localidad = 'CIUDAD'
                    AND l.activo
ON CONFLICT (numero_registro) DO UPDATE SET
  localidad_id = EXCLUDED.localidad_id,
  ruc = EXCLUDED.ruc,
  nombre_comercial = EXCLUDED.nombre_comercial,
  razon_social = EXCLUDED.razon_social,
  actividad = EXCLUDED.actividad,
  clasificacion = EXCLUDED.clasificacion,
  categoria = EXCLUDED.categoria,
  -- La fuente no informa estos campos; no borrar enriquecimientos existentes al repetir el seed.
  direccion = COALESCE(EXCLUDED.direccion, establecimientos_turisticos.direccion),
  telefono = COALESCE(EXCLUDED.telefono, establecimientos_turisticos.telefono),
  latitud = COALESCE(EXCLUDED.latitud, establecimientos_turisticos.latitud),
  longitud = COALESCE(EXCLUDED.longitud, establecimientos_turisticos.longitud),
  activo = TRUE,
  updated_at = CURRENT_TIMESTAMP;

DO $$
DECLARE
  city RECORD;
  selected_count INTEGER;
BEGIN
  IF (SELECT COUNT(*) FROM _catastro_demo_seed) <> 50 THEN
    RAISE EXCEPTION 'El seed de catastro debe contener 50 registros';
  END IF;

  FOR city IN SELECT DISTINCT localidad_nombre FROM _catastro_demo_seed ORDER BY localidad_nombre LOOP
    SELECT COUNT(*) INTO selected_count
    FROM establecimientos_turisticos e
    JOIN _catastro_demo_seed seed ON seed.numero_registro = e.numero_registro
    WHERE seed.localidad_nombre = city.localidad_nombre AND e.activo;
    IF selected_count <> 10 THEN
      RAISE EXCEPTION 'La localidad % debe tener 10 establecimientos activos; tiene %',
        city.localidad_nombre, selected_count;
    END IF;
  END LOOP;

  IF (SELECT COUNT(*) FROM establecimientos_turisticos e
      JOIN _catastro_demo_seed seed ON seed.numero_registro = e.numero_registro
      WHERE e.activo) <> 50 THEN
    RAISE EXCEPTION 'El seed no dejó 50 establecimientos activos';
  END IF;
END;
$$;

COMMIT;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    rows = load_rows(args.source)
    sql = build_sql(args.source, rows)
    args.output.write_text(sql, encoding="utf-8")
    print(f"generated {args.output} with {len(rows)} establishments")


if __name__ == "__main__":
    main()
