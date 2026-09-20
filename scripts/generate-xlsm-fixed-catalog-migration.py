#!/usr/bin/env python3
"""Generate the fixed DPA/classification catalog migration from the XLSM source.

The workbook is an institutional source, not runtime application data.  This script
keeps the generated SQL reproducible while the API and clients continue to consume the
catalogs from PostgreSQL.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


DEFAULT_SOURCE = Path("temp/Centro Cultural Indio Guaranga (2).xlsm")
DEFAULT_OUTPUT = Path("database/migrations/20260920_seed_xlsm_fixed_catalogs.sql")

PROVINCE_DISPLAY_NAMES = {
    "BOLIVAR": "Bolívar",
    "LOSRIOS": "Los Ríos",
    "MANABI": "Manabí",
    "GALAPAGOS": "Galápagos",
    "SUCUMBIOS": "Sucumbíos",
    "SANTODOMINGODELOSTSACHILAS": "Santo Domingo de los Tsáchilas",
    "ZONASENESTUDIO": "Zonas en estudio",
}


def canonical(value: object) -> str:
    text = unicodedata.normalize("NFD", str(value or "").upper())
    text = "".join(char for char in text if unicodedata.category(char) != "Mn")
    return re.sub(r"[^A-Z0-9]+", "", text)


def clean_source_text(value: object) -> str:
    text = str(value or "").replace("_", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"\s*\.\s*$", "", text)


def display_name(value: object) -> str:
    text = clean_source_text(value).lower()
    particles = {"a", "al", "de", "del", "el", "en", "la", "las", "los", "o", "y", "y/o"}
    words: list[str] = []
    for index, word in enumerate(text.split(" ")):
        if index > 0 and word in particles:
            words.append(word)
        elif word:
            words.append(word[:1].upper() + word[1:])
    return " ".join(words)


def sql_literal(value: object) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def sql_values(rows: list[tuple[object, ...]], columns: int) -> str:
    if not rows:
        raise ValueError("SQL VALUES cannot be empty")
    rendered = []
    for row in rows:
        if len(row) != columns:
            raise ValueError(f"Expected {columns} columns, got {len(row)}")
        rendered.append("  (" + ", ".join(sql_literal(value) for value in row) + ")")
    return ",\n".join(rendered)


def read_named_values(workbook, name: str) -> list[object]:
    defined_name = workbook.defined_names.get(name)
    if defined_name is None:
        raise ValueError(f"Named range not found: {name}")

    values: list[object] = []
    for sheet_name, reference in defined_name.destinations:
        sheet = workbook[sheet_name]
        min_column, min_row, max_column, max_row = range_boundaries(reference)
        for row in range(min_row, max_row + 1):
            value = sheet.cell(row=row, column=min_column).value
            if value is not None:
                values.append(value)
    return values


def read_dpa(workbook):
    sheet = workbook["DPA"]
    dpa_rows = []

    provinces: list[tuple[str, str, str]] = []
    for row in range(2, 27):
        raw_name, cod1, cod2 = (sheet.cell(row=row, column=column).value for column in range(1, 4))
        if raw_name is None:
            continue
        code = f"{int(cod1)}{int(cod2)}"
        if len(code) != 2:
            raise ValueError(f"Invalid province code on DPA row {row}: {code}")
        province_key = canonical(raw_name)
        provinces.append(
            (code, PROVINCE_DISPLAY_NAMES.get(province_key, display_name(raw_name)), province_key)
        )

    canton_source_rows: list[tuple[str, str, str]] = []
    canton_row = 29
    province_cantons: list[tuple[str, list[tuple[str, str, str]]]] = []
    for province_code, _, province_key in provinces:
        raw_province_name = next(
            raw_name
            for raw_name, _, raw_key in (
                (sheet.cell(row=row, column=1).value, None, canonical(sheet.cell(row=row, column=1).value))
                for row in range(2, 27)
            )
            if raw_key == province_key
        )
        named_range_name = str(raw_province_name).strip()
        if named_range_name not in workbook.defined_names:
            named_range_name = named_range_name.replace(" ", "_")
        if named_range_name not in workbook.defined_names:
            named_range_name = None
        if named_range_name is None:
            raise ValueError(f"Canton named range not found for province {province_key}")
        canton_names = read_named_values(workbook, named_range_name)
        province_rows: list[tuple[str, str, str]] = []
        for _ in canton_names:
            raw_name, cod1, cod2 = (
                sheet.cell(row=canton_row, column=column).value for column in range(1, 4)
            )
            if raw_name is None:
                raise ValueError(f"Unexpected blank canton row {canton_row}")
            code = f"{int(cod1)}{int(cod2)}"
            if len(code) != 2:
                raise ValueError(f"Invalid canton code on DPA row {canton_row}: {code}")
            province_rows.append((province_code, code, display_name(raw_name)))
            canton_row += 1
        province_cantons.append((province_code, province_rows))
        canton_source_rows.extend(province_rows)

    if canton_row != 256:
        raise ValueError(f"Expected canton section to end at row 255, got {canton_row - 1}")

    parish_rows = []
    for row in range(258, 1508):
        raw_name, cod1, cod2 = (sheet.cell(row=row, column=column).value for column in range(1, 4))
        if raw_name is None:
            raise ValueError(f"Unexpected blank parish row {row}")
        parish_rows.append((display_name(raw_name), f"{int(cod1)}{int(cod2)}"))

    candidate_starts = [
        index
        for index, row in enumerate(range(258, 1508))
        if (sheet.cell(row=row, column=2).value, sheet.cell(row=row, column=3).value)
        in {(0, 1), (5, 0)}
    ]
    # The workbook omits the urban code 01 for Cayambe, so its group starts at
    # COD1/COD2 00/02. Keep this source anomaly explicit and validate it by name.
    cayambe_starts = [
        index
        for index, row in enumerate(range(258, 1508))
        if canonical(sheet.cell(row=row, column=1).value) == "CAYAMBE"
        and (sheet.cell(row=row, column=2).value, sheet.cell(row=row, column=3).value) == (0, 2)
    ]
    if len(cayambe_starts) != 1:
        raise ValueError(f"Expected one Cayambe DPA boundary anomaly, got {len(cayambe_starts)}")
    candidate_starts = sorted(set(candidate_starts + cayambe_starts))
    zone_canton_count = len(province_cantons[-1][1])
    expected_non_zone_count = len(canton_source_rows) - zone_canton_count
    if len(candidate_starts) != expected_non_zone_count:
        raise ValueError(
            "The DPA parish boundary pattern changed: "
            f"found {len(candidate_starts)} starts, expected {expected_non_zone_count}"
        )

    starts = candidate_starts + list(range(len(parish_rows) - zone_canton_count, len(parish_rows)))
    if len(starts) != len(canton_source_rows):
        raise ValueError(f"Expected {len(canton_source_rows)} parish groups, got {len(starts)}")

    parishes: list[tuple[str, str, str, str]] = []
    for group_index, (province_code, canton_code, _) in enumerate(canton_source_rows):
        start = starts[group_index]
        end = starts[group_index + 1] if group_index + 1 < len(starts) else len(parish_rows)
        if start >= end:
            raise ValueError(f"Empty parish group for {province_code}/{canton_code}")
        for parish_name, parish_code in parish_rows[start:end]:
            if len(parish_code) != 2:
                raise ValueError(f"Invalid parish code: {parish_code}")
            parishes.append((province_code, canton_code, parish_code, parish_name))

    if len(provinces) != 25 or len(canton_source_rows) != 227 or len(parishes) != 1250:
        raise ValueError(
            "Unexpected DPA counts: "
            f"{len(provinces)} provinces, {len(canton_source_rows)} cantons, {len(parishes)} parishes"
        )

    return provinces, canton_source_rows, parishes


def read_classification(workbook):
    category_values = read_named_values(workbook, "CATEGORIA")
    if category_values != ["ATRACTIVOS_NATURALES", "MANIFESTACIONES_CULTURALES"]:
        raise ValueError(f"Unexpected category order: {category_values}")
    category_codes = {category_values[0]: "AN", category_values[1]: "MC"}

    sheet = workbook["Clas_AT"]
    type_codes = {}
    for row in range(2, 17):
        name, cod1, cod2 = (sheet.cell(row=row, column=column).value for column in range(1, 4))
        if name is None:
            raise ValueError(f"Unexpected blank type row {row}")
        type_codes[canonical(name)] = f"{int(cod1)}{int(cod2)}"

    types: list[tuple[str, str, str]] = []
    subtypes: list[tuple[str, str, str, str]] = []
    subtype_row = 19
    for category_name in category_values:
        type_names = read_named_values(workbook, category_name)
        category_code = category_codes[category_name]
        for type_name in type_names:
            type_code = type_codes.get(canonical(type_name))
            if type_code is None:
                raise ValueError(f"Type code not found for {type_name}")
            types.append((category_code, type_code, display_name(type_name)))
            subtype_names = read_named_values(workbook, str(type_name))
            for subtype_name in subtype_names:
                raw_name, cod1, cod2 = (
                    sheet.cell(row=subtype_row, column=column).value for column in range(1, 4)
                )
                if raw_name is None:
                    raise ValueError(f"Unexpected blank subtype row {subtype_row}")
                subtype_code = f"{int(cod1)}{int(cod2)}"
                subtypes.append((category_code, type_code, subtype_code, display_name(subtype_name)))
                subtype_row += 1

    if len(types) != 15 or len(subtypes) != 79 or subtype_row != 98:
        raise ValueError(f"Unexpected classification counts: {len(types)} types, {len(subtypes)} subtypes")

    categories = [("AN", "Atractivos naturales"), ("MC", "Manifestaciones culturales")]
    return categories, types, subtypes


def build_sql(source: Path, provinces, cantons, parishes, categories, types, subtypes) -> str:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    lines = [
        "-- Catálogos fijos derivados del XLSM institucional de la ficha turística.",
        f"-- Fuente: {source.as_posix()}",
        f"-- SHA-256: {digest}",
        "-- No incluye localidades ni zonas_turisticas: son catastro operativo separado.",
        "",
        "BEGIN;",
        "",
        "SET LOCAL lock_timeout = '5s';",
        "SET LOCAL statement_timeout = '120s';",
        "",
        "-- DPA: provincias, cantones y parroquias.",
        "INSERT INTO provincias (codigo_dpa, nombre)",
        "VALUES",
        sql_values([(code, name) for code, name, _ in provinces], 2),
        "ON CONFLICT (codigo_dpa) DO UPDATE SET",
        "  nombre = EXCLUDED.nombre,",
        "  activo = TRUE;",
        "",
        "WITH source(codigo_dpa) AS (",
        "  VALUES",
        sql_values([(code,) for code, _, _ in provinces], 1),
        ")",
        "UPDATE provincias AS target",
        "SET activo = FALSE",
        "WHERE NOT EXISTS (SELECT 1 FROM source WHERE source.codigo_dpa = target.codigo_dpa);",
        "",
        "INSERT INTO cantones (provincia_id, codigo_cton, nombre)",
        "SELECT province.id, source.codigo_cton, source.nombre",
        "FROM (VALUES",
        sql_values([(province, code, name) for province, code, name in cantons], 3),
        ") AS source(provincia_codigo, codigo_cton, nombre)",
        "JOIN provincias AS province ON province.codigo_dpa = source.provincia_codigo",
        "ON CONFLICT (provincia_id, codigo_cton) DO UPDATE SET",
        "  nombre = EXCLUDED.nombre,",
        "  activo = TRUE;",
        "",
        "WITH source(provincia_codigo, codigo_cton) AS (",
        "  VALUES",
        sql_values([(province, code) for province, code, _ in cantons], 2),
        ")",
        "UPDATE cantones AS target",
        "SET activo = FALSE",
        "FROM provincias AS province",
        "WHERE target.provincia_id = province.id",
        "  AND NOT EXISTS (",
        "    SELECT 1 FROM source",
        "    WHERE source.provincia_codigo = province.codigo_dpa",
        "      AND source.codigo_cton = target.codigo_cton",
        "  );",
        "",
        "INSERT INTO parroquias (canton_id, codigo_pqa, nombre)",
        "SELECT canton.id, source.codigo_pqa, source.nombre",
        "FROM (VALUES",
        sql_values([(province, canton, code, name) for province, canton, code, name in parishes], 4),
        ") AS source(provincia_codigo, codigo_cton, codigo_pqa, nombre)",
        "JOIN provincias AS province ON province.codigo_dpa = source.provincia_codigo",
        "JOIN cantones AS canton",
        "  ON canton.provincia_id = province.id",
        " AND canton.codigo_cton = source.codigo_cton",
        "ON CONFLICT (canton_id, codigo_pqa) DO UPDATE SET",
        "  nombre = EXCLUDED.nombre,",
        "  activo = TRUE;",
        "",
        "WITH source(provincia_codigo, codigo_cton, codigo_pqa) AS (",
        "  VALUES",
        sql_values([(province, canton, code) for province, canton, code, _ in parishes], 3),
        ")",
        "UPDATE parroquias AS target",
        "SET activo = FALSE",
        "FROM cantones AS canton",
        "JOIN provincias AS province ON province.id = canton.provincia_id",
        "WHERE target.canton_id = canton.id",
        "  AND NOT EXISTS (",
        "    SELECT 1 FROM source",
        "    WHERE source.provincia_codigo = province.codigo_dpa",
        "      AND source.codigo_cton = canton.codigo_cton",
        "      AND source.codigo_pqa = target.codigo_pqa",
        "  );",
        "",
        "-- Clasificación oficial de atractivos naturales y manifestaciones culturales.",
        "INSERT INTO categorias_atractivo (codigo, nombre)",
        "VALUES",
        sql_values(categories, 2),
        "ON CONFLICT (codigo) DO UPDATE SET",
        "  nombre = EXCLUDED.nombre,",
        "  activo = TRUE;",
        "",
        "WITH source(categoria_codigo, codigo, nombre) AS (",
        "  VALUES",
        sql_values(types, 3),
        ")",
        "INSERT INTO tipos_atractivo (categoria_id, codigo, nombre)",
        "SELECT category.id, source.codigo, source.nombre",
        "FROM source",
        "JOIN categorias_atractivo AS category ON category.codigo = source.categoria_codigo",
        "ON CONFLICT (categoria_id, codigo) DO UPDATE SET",
        "  nombre = EXCLUDED.nombre,",
        "  activo = TRUE;",
        "",
        "WITH source(categoria_codigo, tipo_codigo, codigo, nombre) AS (",
        "  VALUES",
        sql_values(subtypes, 4),
        ")",
        "INSERT INTO subtipos_atractivo (tipo_atractivo_id, codigo, nombre)",
        "SELECT type.id, source.codigo, source.nombre",
        "FROM source",
        "JOIN categorias_atractivo AS category ON category.codigo = source.categoria_codigo",
        "JOIN tipos_atractivo AS type",
        "  ON type.categoria_id = category.id",
        " AND type.codigo = source.tipo_codigo",
        "ON CONFLICT (tipo_atractivo_id, codigo) DO UPDATE SET",
        "  nombre = EXCLUDED.nombre,",
        "  activo = TRUE;",
        "",
        "WITH source(categoria_codigo, tipo_codigo, codigo) AS (",
        "  VALUES",
        sql_values([(category, type_code, subtype_code) for category, type_code, subtype_code, _ in subtypes], 3),
        ")",
        "UPDATE subtipos_atractivo AS target",
        "SET activo = FALSE",
        "FROM tipos_atractivo AS type",
        "JOIN categorias_atractivo AS category ON category.id = type.categoria_id",
        "WHERE target.tipo_atractivo_id = type.id",
        "  AND NOT EXISTS (",
        "    SELECT 1 FROM source",
        "    WHERE source.categoria_codigo = category.codigo",
        "      AND source.tipo_codigo = type.codigo",
        "      AND source.codigo = target.codigo",
        "  );",
        "",
        "WITH source(categoria_codigo, codigo) AS (",
        "  VALUES",
        sql_values([(category, type_code) for category, type_code, _ in types], 2),
        ")",
        "UPDATE tipos_atractivo AS target",
        "SET activo = FALSE",
        "FROM categorias_atractivo AS category",
        "WHERE target.categoria_id = category.id",
        "  AND NOT EXISTS (",
        "    SELECT 1 FROM source",
        "    WHERE source.categoria_codigo = category.codigo",
        "      AND source.codigo = target.codigo",
        "  );",
        "",
        "WITH source(codigo) AS (",
        "  VALUES",
        sql_values([(code,) for code, _ in categories], 1),
        ")",
        "UPDATE categorias_atractivo AS target",
        "SET activo = FALSE",
        "WHERE NOT EXISTS (SELECT 1 FROM source WHERE source.codigo = target.codigo);",
        "",
        "DO $$",
        "DECLARE",
        "  actual BIGINT;",
        "BEGIN",
        "  SELECT count(*) INTO actual FROM provincias WHERE activo;",
        "  IF actual <> 25 THEN RAISE EXCEPTION 'DPA: se esperaban 25 provincias activas, hay %', actual; END IF;",
        "  SELECT count(*) INTO actual FROM cantones WHERE activo;",
        "  IF actual <> 227 THEN RAISE EXCEPTION 'DPA: se esperaban 227 cantones activos, hay %', actual; END IF;",
        "  SELECT count(*) INTO actual FROM parroquias WHERE activo;",
        "  IF actual <> 1250 THEN RAISE EXCEPTION 'DPA: se esperaban 1250 parroquias activas, hay %', actual; END IF;",
        "  SELECT count(*) INTO actual FROM categorias_atractivo WHERE activo;",
        "  IF actual <> 2 THEN RAISE EXCEPTION 'Clasificación: se esperaban 2 categorías activas, hay %', actual; END IF;",
        "  SELECT count(*) INTO actual FROM tipos_atractivo WHERE activo;",
        "  IF actual <> 15 THEN RAISE EXCEPTION 'Clasificación: se esperaban 15 tipos activos, hay %', actual; END IF;",
        "  SELECT count(*) INTO actual FROM subtipos_atractivo WHERE activo;",
        "  IF actual <> 79 THEN RAISE EXCEPTION 'Clasificación: se esperaban 79 subtipos activos, hay %', actual; END IF;",
        "END $$;",
        "",
        "COMMIT;",
        "",
    ]
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    if not args.source.exists():
        raise SystemExit(f"Source workbook not found: {args.source}")
    workbook = load_workbook(args.source, data_only=False, read_only=False, keep_vba=True)
    provinces, cantons, parishes = read_dpa(workbook)
    categories, types, subtypes = read_classification(workbook)
    sql = build_sql(args.source, provinces, cantons, parishes, categories, types, subtypes)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(sql, encoding="utf-8")
    print(
        f"Generated {args.output}: {len(provinces)} provinces, {len(cantons)} cantons, "
        f"{len(parishes)} parishes, {len(categories)} categories, {len(types)} types, "
        f"{len(subtypes)} subtypes"
    )


if __name__ == "__main__":
    main()
