#!/usr/bin/env python3
"""Generate the hierarchical establishment taxonomy from the national XLSX."""

from __future__ import annotations

import argparse
import hashlib
import re
import unicodedata
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SOURCE = Path("temp/Consolidado-Nacional-2026-publico-8 (1).xlsx")
DEFAULT_OUTPUT = Path("database/migrations/20260921_aa_establishment_taxonomy.sql")


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def fold(value: str) -> str:
    normalized = unicodedata.normalize("NFD", value.upper())
    without_marks = "".join(
        char for char in normalized if unicodedata.category(char) != "Mn"
    )
    return re.sub(r"\s+", " ", without_marks).strip()


def slug(value: str) -> str:
    return re.sub(r"[^A-Z0-9]+", "_", fold(value)).strip("_")


def sql(value: str | None) -> str:
    if value is None:
        return "NULL"
    return "'" + value.replace("'", "''") + "'"


def canonical_classification(value: str) -> str:
    result = re.sub(r"\s+", " ", value.strip())
    result = result.replace(
        "GUÍA DE TURISMO ESPECIALIZADO EN AVENTURA",
        "GUÍA DE TURISMO DE AVENTURA",
    )
    result = result.replace(
        "GUÍA DE TURISMO NACIONAL ESPCIALIZADO EN PATRIMON IO TURÍSTICO",
        "GUÍA DE TURISMO NACIONAL ESPECIALIZADO EN PATRIMONIO TURÍSTICO",
    )
    result = re.sub(
        r"GUÍA DE TURISMO NACIONAL ESPECIALIZADO EN PATRIMONIO(?! TURÍSTICO)",
        "GUÍA DE TURISMO NACIONAL ESPECIALIZADO EN PATRIMONIO TURÍSTICO",
        result,
    )
    return result


def canonical_category(value: str) -> str:
    result = re.sub(r"\s+", " ", value.strip())
    key = fold(result)
    match = re.match(
        r"^\(?([1-9])\)?\s+(UN|UNA|DOS|TRES|CUATRO|CINCO)\s+(TAZAS?|COPAS?|TENEDOR(?:ES)?)$",
        key,
    )
    if match:
        number, _number_word, unit = match.groups()
        unit = "TENEDOR" if unit.startswith("TENEDOR") else unit.rstrip("S")
        singular = {"TAZA": "Taza", "COPA": "Copa", "TENEDOR": "Tenedor"}[unit]
        plural = {"TAZA": "Tazas", "COPA": "Copas", "TENEDOR": "Tenedores"}[unit]
        return f"{number} {singular if number == '1' else plural}"
    match = re.match(r"^([1-9])\s+ESTRELLAS?$", key)
    if match:
        number = match.group(1)
        return f"{number} Estrella{'' if number == '1' else 's'}"
    if key in {"CATEGORIA UNICA", "CATEGORIAUNICA"}:
        return "Categoría única"
    if key == "CATEGORIA UNO":
        return "Categoría Uno"
    if key == "CATEGORIA DOS":
        return "Categoría Dos"
    known = {
        "CLASE TURISTA": "Clase turista",
        "PRIMERA CLASE": "Primera clase",
        "CLASE ECONOMICA": "Clase económica",
        "LUJO": "Lujo",
        "PRIMERA": "Primera",
        "SEGUNDA": "Segunda",
    }
    return known.get(key, result)


def stable_code(prefix: str, *values: str) -> str:
    digest = hashlib.sha1("|".join(values).encode("utf-8")).hexdigest()[:14]
    return f"{prefix}_{digest}"[:120]


def load_taxonomy(source: Path):
    worksheet = load_workbook(source, read_only=True, data_only=True).active
    rows = worksheet.iter_rows(values_only=True)
    headers = tuple(text(value) for value in next(rows))
    required = ("Actividad / Modalidad", "Clasificación", "Categoría")
    if any(field not in headers for field in required):
        raise ValueError(f"Missing required XLSX headers: {headers!r}")
    index = {header: position for position, header in enumerate(headers)}

    activities: OrderedDict[str, str] = OrderedDict()
    classifications: OrderedDict[tuple[str, str], dict[str, object]] = OrderedDict()
    categories: OrderedDict[tuple[str, str, str], dict[str, object]] = OrderedDict()

    for raw in rows:
        activity = text(raw[index["Actividad / Modalidad"]])
        source_classification = text(raw[index["Clasificación"]])
        source_category = text(raw[index["Categoría"]])
        if not activity or not source_classification or not source_category:
            continue
        classification = canonical_classification(source_classification)
        category = canonical_category(source_category)
        activities.setdefault(activity, activity)
        classification_key = (activity, classification)
        classification_entry = classifications.setdefault(
            classification_key,
            {"activity": activity, "name": classification, "aliases": OrderedDict()},
        )
        aliases = classification_entry["aliases"]
        assert isinstance(aliases, OrderedDict)
        aliases[source_classification] = source_classification
        category_key = (activity, classification, category)
        category_entry = categories.setdefault(
            category_key,
            {
                "activity": activity,
                "classification": classification,
                "name": category,
                "aliases": OrderedDict(),
            },
        )
        category_aliases = category_entry["aliases"]
        assert isinstance(category_aliases, OrderedDict)
        category_aliases[source_category] = source_category

    return activities, classifications, categories


def render_values(rows: list[tuple[str, ...]]) -> str:
    return ",\n".join("  (" + ", ".join(sql(value) for value in row) + ")" for row in rows)


def build_sql(source: Path) -> str:
    activities, classifications, categories = load_taxonomy(source)
    activity_rows = [
        (stable_code("ACT", activity), activity) for activity in activities.values()
    ]
    classification_rows = [
        (
            stable_code("CLS", activity, str(entry["name"])),
            activity,
            str(entry["name"]),
        )
        for (activity, _classification), entry in classifications.items()
    ]
    category_rows = [
        (
            stable_code("CAT", activity, classification, name),
            activity,
            classification,
            name,
            str(index),
        )
        for index, ((activity, classification, name), _entry) in enumerate(
            categories.items(), start=1
        )
    ]
    classification_alias_rows = [
        (stable_code("CLS", activity, str(entry["name"])), str(alias))
        for (activity, _name), entry in classifications.items()
        for alias in entry["aliases"]
    ]
    category_alias_rows = [
        (stable_code("CAT", activity, classification, name), str(alias))
        for (activity, classification, name), entry in categories.items()
        for alias in entry["aliases"]
    ]

    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    return f"""-- Catálogo jerárquico del catastro derivado del consolidado nacional.
-- Fuente: {source.as_posix()}
-- SHA-256: {digest}
-- Las columnas de texto del catastro se conservan; los aliases permiten resolver
-- variantes del archivo sin perder la trazabilidad del valor de origen.

BEGIN;

SET LOCAL lock_timeout = '5s';
SET LOCAL statement_timeout = '120s';

CREATE TABLE IF NOT EXISTS catalogo_catastro_actividades (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    codigo VARCHAR(120) NOT NULL UNIQUE,
    nombre VARCHAR(180) NOT NULL UNIQUE,
    activo BOOLEAN NOT NULL DEFAULT TRUE
);

CREATE TABLE IF NOT EXISTS catalogo_catastro_clasificaciones (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    actividad_id BIGINT NOT NULL REFERENCES catalogo_catastro_actividades(id) ON UPDATE CASCADE ON DELETE RESTRICT,
    codigo VARCHAR(120) NOT NULL UNIQUE,
    nombre VARCHAR(180) NOT NULL,
    activo BOOLEAN NOT NULL DEFAULT TRUE,
    UNIQUE (actividad_id, nombre)
);

CREATE INDEX IF NOT EXISTS idx_catastro_clasificaciones_actividad
    ON catalogo_catastro_clasificaciones (actividad_id, activo, nombre);

CREATE TABLE IF NOT EXISTS catalogo_catastro_categorias (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    clasificacion_id BIGINT NOT NULL REFERENCES catalogo_catastro_clasificaciones(id) ON UPDATE CASCADE ON DELETE RESTRICT,
    codigo VARCHAR(120) NOT NULL UNIQUE,
    nombre VARCHAR(120) NOT NULL,
    orden SMALLINT NOT NULL DEFAULT 0,
    activo BOOLEAN NOT NULL DEFAULT TRUE,
    UNIQUE (clasificacion_id, nombre)
);

CREATE INDEX IF NOT EXISTS idx_catastro_categorias_clasificacion
    ON catalogo_catastro_categorias (clasificacion_id, activo, orden, nombre);

CREATE TABLE IF NOT EXISTS catalogo_catastro_clasificacion_aliases (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    clasificacion_id BIGINT NOT NULL REFERENCES catalogo_catastro_clasificaciones(id) ON UPDATE CASCADE ON DELETE CASCADE,
    valor_fuente VARCHAR(180) NOT NULL,
    UNIQUE (clasificacion_id, valor_fuente)
);

CREATE TABLE IF NOT EXISTS catalogo_catastro_categoria_aliases (
    id BIGINT GENERATED ALWAYS AS IDENTITY PRIMARY KEY,
    categoria_id BIGINT NOT NULL REFERENCES catalogo_catastro_categorias(id) ON UPDATE CASCADE ON DELETE CASCADE,
    valor_fuente VARCHAR(120) NOT NULL,
    UNIQUE (categoria_id, valor_fuente)
);

ALTER TABLE establecimientos_turisticos
    ADD COLUMN IF NOT EXISTS actividad_catalogo_id BIGINT,
    ADD COLUMN IF NOT EXISTS clasificacion_catalogo_id BIGINT,
    ADD COLUMN IF NOT EXISTS categoria_catalogo_id BIGINT;

DO $$
BEGIN
  IF NOT EXISTS (SELECT 1 FROM pg_constraint WHERE conname = 'fk_establecimientos_actividad_catalogo') THEN
    ALTER TABLE establecimientos_turisticos
      ADD CONSTRAINT fk_establecimientos_actividad_catalogo
      FOREIGN KEY (actividad_catalogo_id) REFERENCES catalogo_catastro_actividades(id)
      ON UPDATE CASCADE ON DELETE RESTRICT;
  END IF;
  IF NOT EXISTS (SELECT 1 FROM pg_constraint WHERE conname = 'fk_establecimientos_clasificacion_catalogo') THEN
    ALTER TABLE establecimientos_turisticos
      ADD CONSTRAINT fk_establecimientos_clasificacion_catalogo
      FOREIGN KEY (clasificacion_catalogo_id) REFERENCES catalogo_catastro_clasificaciones(id)
      ON UPDATE CASCADE ON DELETE RESTRICT;
  END IF;
  IF NOT EXISTS (SELECT 1 FROM pg_constraint WHERE conname = 'fk_establecimientos_categoria_catalogo') THEN
    ALTER TABLE establecimientos_turisticos
      ADD CONSTRAINT fk_establecimientos_categoria_catalogo
      FOREIGN KEY (categoria_catalogo_id) REFERENCES catalogo_catastro_categorias(id)
      ON UPDATE CASCADE ON DELETE RESTRICT;
  END IF;
END;
$$;

CREATE INDEX IF NOT EXISTS idx_establecimientos_actividad_catalogo
    ON establecimientos_turisticos (actividad_catalogo_id);
CREATE INDEX IF NOT EXISTS idx_establecimientos_clasificacion_catalogo
    ON establecimientos_turisticos (clasificacion_catalogo_id);
CREATE INDEX IF NOT EXISTS idx_establecimientos_categoria_catalogo
    ON establecimientos_turisticos (categoria_catalogo_id);

INSERT INTO catalogo_catastro_actividades (codigo, nombre)
VALUES
{render_values(activity_rows)}
ON CONFLICT (codigo) DO UPDATE SET nombre = EXCLUDED.nombre, activo = TRUE;

WITH source(codigo, actividad, nombre) AS (
  VALUES
{render_values(classification_rows)}
)
INSERT INTO catalogo_catastro_clasificaciones (actividad_id, codigo, nombre)
SELECT activity.id, source.codigo, source.nombre
FROM source
JOIN catalogo_catastro_actividades activity ON activity.nombre = source.actividad
ON CONFLICT (codigo) DO UPDATE SET
  actividad_id = EXCLUDED.actividad_id,
  nombre = EXCLUDED.nombre,
  activo = TRUE;

WITH source(codigo, actividad, clasificacion, nombre, orden) AS (
  VALUES
{render_values(category_rows)}
)
INSERT INTO catalogo_catastro_categorias (clasificacion_id, codigo, nombre, orden)
SELECT classification.id, source.codigo, source.nombre, source.orden::SMALLINT
FROM source
JOIN catalogo_catastro_actividades activity ON activity.nombre = source.actividad
JOIN catalogo_catastro_clasificaciones classification
  ON classification.actividad_id = activity.id AND classification.nombre = source.clasificacion
ON CONFLICT (codigo) DO UPDATE SET
  clasificacion_id = EXCLUDED.clasificacion_id,
  nombre = EXCLUDED.nombre,
  orden = EXCLUDED.orden,
  activo = TRUE;

WITH source(codigo, valor_fuente) AS (
  VALUES
{render_values(classification_alias_rows)}
)
INSERT INTO catalogo_catastro_clasificacion_aliases (clasificacion_id, valor_fuente)
SELECT classification.id, source.valor_fuente
FROM source
JOIN catalogo_catastro_clasificaciones classification ON classification.codigo = source.codigo
ON CONFLICT (clasificacion_id, valor_fuente) DO NOTHING;

WITH source(codigo, valor_fuente) AS (
  VALUES
{render_values(category_alias_rows)}
)
INSERT INTO catalogo_catastro_categoria_aliases (categoria_id, valor_fuente)
SELECT category.id, source.valor_fuente
FROM source
JOIN catalogo_catastro_categorias category ON category.codigo = source.codigo
ON CONFLICT (categoria_id, valor_fuente) DO NOTHING;

UPDATE establecimientos_turisticos establishment
SET actividad_catalogo_id = activity.id,
    clasificacion_catalogo_id = classification.id
FROM catalogo_catastro_actividades activity
JOIN catalogo_catastro_clasificaciones classification
  ON classification.actividad_id = activity.id
JOIN catalogo_catastro_clasificacion_aliases alias
  ON alias.clasificacion_id = classification.id
WHERE activity.nombre = establishment.actividad
  AND alias.valor_fuente = establishment.clasificacion;

UPDATE establecimientos_turisticos establishment
SET categoria_catalogo_id = category.id
FROM catalogo_catastro_categorias category
JOIN catalogo_catastro_categoria_aliases alias ON alias.categoria_id = category.id
WHERE category.clasificacion_id = establishment.clasificacion_catalogo_id
  AND alias.valor_fuente = establishment.categoria;

ALTER TABLE auditoria_catalogos
  DROP CONSTRAINT IF EXISTS auditoria_catalogos_catalogo_codigo_check;

ALTER TABLE auditoria_catalogos
  ADD CONSTRAINT auditoria_catalogos_catalogo_codigo_check
  CHECK (catalogo_codigo IN (
    'ACCESSIBILITY', 'ACTIVITY', 'FACILITY', 'ESTABLISHMENT', 'ESTABLISHMENT_CATEGORY'
  ));

COMMIT;
"""


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    args.output.write_text(build_sql(args.source), encoding="utf-8")


if __name__ == "__main__":
    main()
