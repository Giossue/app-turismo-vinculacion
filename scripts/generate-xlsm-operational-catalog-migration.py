#!/usr/bin/env python3
"""Generate the operational catalog seed from the institutional XLSM.

The workbook is the source for the labels and order of the fixed choices used by the
administrative center record.  The generated SQL is deliberately data-only: schema
changes belong to the versioned migrations that created the catalog tables.
"""

from __future__ import annotations

import argparse
import hashlib
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


DEFAULT_SOURCE = Path("temp/Centro Cultural Indio Guaranga (2).xlsm")
DEFAULT_OUTPUT = Path("database/migrations/20260920_seed_xlsm_operational_catalogs.sql")


def clean_text(value: object) -> str:
    """Collapse workbook whitespace without changing the institutional wording."""

    return re.sub(r"\s+", " ", str(value or "").replace("_", " ")).strip()


def choice(sheet, coordinate: str) -> str:
    value = clean_text(sheet[coordinate].value)
    value = re.sub(r"^[a-z]\.\s*", "", value, flags=re.IGNORECASE)
    value = re.sub(r"^[a-z]\s+(?=[A-ZÁÉÍÓÚÑ])", "", value, flags=re.IGNORECASE)
    return value.strip()


def sql_literal(value: object) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def sql_values(rows: list[tuple[object, ...]], columns: int) -> str:
    if not rows:
        raise ValueError("SQL VALUES cannot be empty")
    rendered: list[str] = []
    for row in rows:
        if len(row) != columns:
            raise ValueError(f"Expected {columns} columns, got {len(row)}")
        rendered.append("  (" + ", ".join(sql_literal(value) for value in row) + ")")
    return ",\n".join(rendered)


def read_named_values(workbook, name: str) -> list[object]:
    defined_name = workbook.defined_names.get(name)
    if defined_name is None:
        raise ValueError(f"Named range not found: {name}")
    values: list[object] = []
    for sheet_name, reference in defined_name.destinations:
        sheet = workbook[sheet_name]
        min_column, min_row, max_column, max_row = range_boundaries(reference)
        for row in range(min_row, max_row + 1):
            value = sheet.cell(row=row, column=min_column).value
            if value is not None:
                values.append(value)
    return values


def slug(value: str) -> str:
    normalized = unicodedata.normalize("NFD", value.upper())
    normalized = "".join(char for char in normalized if unicodedata.category(char) != "Mn")
    return "_".join(re.findall(r"[A-Z0-9]+", normalized))


def bounded_code(prefix: str, value: str, index: int, max_length: int) -> str:
    candidate = f"{prefix}_{slug(value)}"
    if len(candidate) <= max_length:
        return candidate
    return f"{prefix}_{index:02d}"


def assert_count(label: str, rows: list[object], expected: int) -> None:
    if len(rows) != expected:
        raise ValueError(f"Unexpected {label} count: got {len(rows)}, expected {expected}")


def extract_accessibility(workbook) -> list[tuple[str, str, str, int]]:
    sheet = workbook["ficha_Accesibilidad"]
    groups = [
        ("GENERAL", 7, 22),
        ("FISICA", 22, 37),
        ("VISUAL", 37, 50),
        ("AUDITIVA", 50, 58),
        ("COGNITIVA", 58, 63),
    ]
    criteria: list[tuple[str, str, str, int]] = []
    for type_code, start, end in groups:
        values = [clean_text(sheet.cell(row=row, column=1).value) for row in range(start + 1, end)]
        values = [value for value in values if value]
        for order, description in enumerate(values, start=1):
            criteria.append((type_code, f"{type_code}_{order:02d}", description, order))
    assert_count("accessibility criteria", criteria, 51)
    expected_counts = {"GENERAL": 14, "FISICA": 14, "VISUAL": 12, "AUDITIVA": 7, "COGNITIVA": 4}
    for type_code, expected in expected_counts.items():
        actual = sum(1 for row in criteria if row[0] == type_code)
        if actual != expected:
            raise ValueError(f"Unexpected {type_code} accessibility criteria count: {actual}")
    return criteria


def extract_operational_catalogs(workbook) -> dict[str, list[tuple[object, ...]]]:
    sheet = workbook["Ficha_Jerarquia"]
    values_sheet = workbook["Valores"]

    transport_coordinates = [
        "B55", "E55", "H55", "L55", "P55", "S55",
        "B56", "E56", "H56", "L56", "P56", "S56", "B57", "E57",
    ]
    transport_codes = [
        "BUS", "BUSETA", "TRANSPORTE_4X4", "TAXI", "MOTOTAXI", "TELEFERICO",
        "LANCHA", "BOTE", "BARCO", "CANOA", "AVION", "AVIONETA", "HELICOPTERO", "OTRO",
    ]
    transport_names = [choice(sheet, coordinate) for coordinate in transport_coordinates]
    assert_count("transport types", transport_names, 14)
    transports = list(zip(transport_codes, transport_names))

    frequency_coordinates = ["K61", "L61", "M61", "N61"]
    frequency_codes = ["DIARIO", "SEMANAL", "MENSUAL", "EVENTUAL"]
    frequencies = list(zip(frequency_codes, [choice(sheet, coordinate) for coordinate in frequency_coordinates]))

    criteria = extract_accessibility(workbook)

    plant: list[tuple[object, ...]] = []
    for group, prefix, coordinates, units in [
        (
            "ALOJAMIENTO",
            "ALOJAMIENTO",
            ["B77", "B78", "B79", "B80", "B81", "B82", "B83", "B84", "B85"],
            ("Establecimientos registrados", "Número de Habitaciones", "Número de Plazas"),
        ),
        (
            "ALIMENTOS_BEBIDAS",
            "ALIMENTOS",
            ["B88", "B89", "B90", "B91"],
            ("Establecimientos registrados", "Número de Mesas", "Número de Plazas"),
        ),
        (
            "AGENCIAS_VIAJE",
            "AGENCIA",
            ["B94", "B95", "B96"],
            ("Establecimientos registrados", None, None),
        ),
    ]:
        for index, coordinate in enumerate(coordinates, start=1):
            name = choice(sheet, coordinate)
            plant.append((group, f"{prefix}_{slug(name)}", name, *units))
    guide_names = ["Local", "Nacional", "Nacional especializado en cultura", "Nacional especializado en aventura"]
    for name in guide_names:
        plant.append(("GUIAS", f"GUIA_{slug(name)}", name, "Personas", None, None))
    assert_count("tourism plant types", plant, 20)

    facility_categories = [
        ("GESTION", choice(sheet, "B104")),
        ("OBSERVACION", choice(sheet, "B109")),
        ("RECORRIDO", choice(sheet, "B113")),
        ("SERVICIO", choice(sheet, "B117")),
        ("OTROS", choice(sheet, "B119")),
    ]
    facility_types: list[tuple[object, ...]] = []
    facility_code_aliases = {
        "Punto de Información": "PUNTO_INFORMACION",
        "I-Tur": "I_TUR",
        "Centro de interpretación": "CENTRO_INTERPRETACION",
        "Centro de facilitación turística": "CENTRO_FACILITACION",
        "Centro de recepción de visitantes": "CENTRO_RECEPCION",
        "Garitas de guardianía": "GARITA",
        "Miradores": "MIRADOR",
        "Torres de avistamiento de aves": "TORRE_AVISTAMIENTO",
        "Torres de vigilancia para salvavidas": "TORRE_SALVAVIDAS",
        "Senderos": "SENDERO",
        "Estaciones de sombra y descanso": "SOMBRA_DESCANSO",
        "Áreas de acampar": "AREA_ACAMPAR",
        "Refugio de alta montaña": "REFUGIO_MONTAÑA",
        "Baterías sanitarias": "BATERIA_SANITARIA",
        "Estacionamientos": "ESTACIONAMIENTO",
    }
    for category_code, coordinates in [
        ("GESTION", ["E104", "E105", "E106", "E107", "E108"]),
        ("OBSERVACION", ["E109", "E110", "E111", "E112"]),
        ("RECORRIDO", ["E113", "E114", "E115", "E116"]),
        ("SERVICIO", ["E117", "E118"]),
    ]:
        for coordinate in coordinates:
            name = choice(sheet, coordinate)
            facility_types.append((category_code, facility_code_aliases[name], name))
    facility_types.append(("OTROS", "OTRO", choice(sheet, "B119")))
    assert_count("facility types", facility_types, 16)

    complementary_coordinates = ["B123", "H123", "B124", "F124", "J124"]
    complementary_codes = [
        "ALQUILER_VENTA_EQUIPO", "VENTA_ARTESANIAS_MERCH", "CASA_CAMBIO", "CAJERO_AUTOMATICO", "OTRO",
    ]
    complementary = list(zip(complementary_codes, [choice(sheet, coordinate) for coordinate in complementary_coordinates]))

    conservation_coordinates = ["B129", "G129", "L129", "R129"]
    conservation_codes = ["CONSERVADO", "ALTERADO", "EN_PROCESO_DE_DETERIORO", "DETERIORADO"]
    conservation_states = list(zip(conservation_codes, [choice(sheet, coordinate) for coordinate in conservation_coordinates]))

    natural_factor_coordinates = ["B133", "B134", "B135", "B136", "B137", "B138"]
    anthropic_factor_coordinates = [
        "H133", "M133", "R133", "H134", "M134", "R134", "H135", "M135", "R135",
        "H136", "M136", "R136", "H137", "M137", "R137", "B138",
    ]
    factors: list[tuple[object, ...]] = []
    for origin, coordinates in [("NATURAL", natural_factor_coordinates), ("ANTROPICO", anthropic_factor_coordinates)]:
        for coordinate in coordinates:
            name = choice(sheet, coordinate)
            factors.append((origin, bounded_code(origin[:3], name, len(factors) + 1, 50), name))
    assert_count("alteration factors", factors, 22)

    basic_categories = [
        ("AGUA", "Agua", "Agua"),
        ("ENERGIA", "Energía", "Energía"),
        ("SANEAMIENTO", "Saneamiento", "Saneamiento"),
        ("DESECHOS", "Disposición de desechos", "Desechos"),
    ]
    basic_types: list[tuple[object, ...]] = []
    for category_code, category_name, named_range in basic_categories:
        for index, raw_name in enumerate(read_named_values(workbook, named_range), start=1):
            name = clean_text(raw_name)
            basic_types.append((category_code, bounded_code(category_code, name, index, 40), name))
    assert_count("basic service types", basic_types, 23)

    signage: list[tuple[object, ...]] = []
    for environment, coordinates in [
        ("URBANA", [f"D{row}" for row in range(170, 178)]),
        ("NATURAL", [f"D{row}" for row in range(178, 190)]),
        ("INFORMATIVA", ["D190", "D191"]),
        ("SEGURIDAD", ["D192"]),
    ]:
        for index, coordinate in enumerate(coordinates, start=1):
            signage.append((environment, f"{environment}_{index:02d}", choice(sheet, coordinate)))
    signage.append(("OTROS", "OTRO", choice(sheet, "B193")))
    assert_count("signage types", signage, 24)

    material_coordinates = ["J169", "L169", "N169"]
    material_codes = ["MADERA", "ALUMINIO", "OTRO"]
    signage_materials = list(zip(material_codes, [choice(sheet, coordinate) for coordinate in material_coordinates]))

    health_coordinates = ["B197", "B198", "B199", "B200", "B201"]
    health_codes = ["HOSPITAL_CLINICA", "CENTRO_SALUD", "DISPENSARIO_MEDICO", "BOTIQUIN_PRIMEROS_AUXILIOS", "OTRO"]
    health_services = list(zip(health_codes, [choice(sheet, coordinate) for coordinate in health_coordinates]))

    security_coordinates = ["B204", "B205", "B206", "B207"]
    security_codes = ["PRIVADA", "POLICIA_NACIONAL", "POLICIA_MUNICIPAL", "OTRO"]
    security_services = list(zip(security_codes, [choice(sheet, coordinate) for coordinate in security_coordinates]))

    communication: list[tuple[object, ...]] = []
    for group, coordinates in [
        ("TELEFONIA", ["B212", "B213", "B214"]),
        ("INTERNET", ["E212", "I212", "E213", "I213", "E214"]),
    ]:
        for index, coordinate in enumerate(coordinates, start=1):
            name = choice(sheet, coordinate)
            communication.append((group, bounded_code(group, name, index, 40), name))
    assert_count("communication types", communication, 8)

    threat_coordinates = ["B220", "G220", "M220", "R220", "B221", "G221", "M221", "R221"]
    threats = [(bounded_code("AMENAZA", choice(sheet, coordinate), index, 30), choice(sheet, coordinate))
               for index, coordinate in enumerate(threat_coordinates, start=1)]

    policy_questions = [
        ("PLAN_DESARROLLO_GAD", "¿El atractivo se encuentra dentro del plan de desarrollo turístico del GAD?", 1),
        ("PLANIFICACION_TERRITORIAL", clean_text(sheet["B226"].value), 2),
        ("REGULACIONES_APLICABLES", clean_text(sheet["B228"].value), 3),
        ("ORDENANZAS_APLICABLES", clean_text(sheet["B230"].value), 4),
    ]

    activity_cells = {
        "AGUA": [
            "B236", "F236", "J236", "N236", "R236", "B237", "F237", "J237", "N237", "R237",
            "B238", "F238", "J238", "N238", "R238", "B239", "F239", "J239", "N239", "R239",
        ],
        "AIRE": ["B242", "F242", "J242", "N242"],
        "TIERRA": [
            "B245", "F245", "J245", "N245", "R245", "B246", "F246", "J246", "N246", "R246",
            "B247", "F247", "J247", "N247",
        ],
        "CULTURA": [
            "B251", "I251", "P251", "B252", "I252", "P252", "B253", "I253", "P253",
            "B254", "I254", "P254", "B255", "I255", "P255", "B256", "I256", "P256",
        ],
    }
    existing_activity_codes = {
        "Buceo": "BUCEO", "Rafting": "RAFTING", "Pesca deportiva": "PESCA_DEPORTIVA", "Canopy": "CANOPY",
        "Parapente": "PARAPENTE", "Alas Delta": "ALA_DELTA", "Montañismo": "MONTANISMO", "Escalada": "ESCALADA",
        "Senderismo": "SENDERISMO", "Cicloturismo": "CICLOTURISMO", "Cabalgata": "CABALGATA", "Caminata": "CAMINATA",
        "Camping": "CAMPING", "Picnic": "PICNIC", "Observación de flora y fauna": "OBSERVACION_FLORA_FAUNA",
        "Recorridos guiados": "RECORRIDO_GUIADO", "Fotografía": "FOTOGRAFIA", "Degustación de platos tradicionales": "DEGUSTACION",
        "Participación de la celebración": "CELEBRACIONES", "Compra de artesanías": "ARTESANIAS", "Medicina ancestral": "MEDICINA_ANCESTRAL",
    }
    activities: list[tuple[object, ...]] = []
    for group, coordinates in activity_cells.items():
        for index, coordinate in enumerate(coordinates, start=1):
            name = choice(sheet, coordinate)
            code = existing_activity_codes.get(name, bounded_code(group, name, index, 50))
            activities.append((group, code, name))
    assert_count("tourism activities", activities, 56)

    promotion_coordinates = ["B263", "B264", "B265", "B266", "B267", "B268", "B269", "B270"]
    promotion_codes = [
        "PAGINA_WEB", "RED_SOCIAL", "REVISTA_ESPECIALIZADA", "MATERIAL_POP",
        "OFICINA_INFORMACION_TURISTICA", "MEDIOS_COMUNICACION", "FERIAS_TURISTICAS", "OTRO",
    ]
    promotion_media = list(zip(promotion_codes, [choice(sheet, coordinate) for coordinate in promotion_coordinates]))

    training: list[tuple[object, ...]] = []
    for group, coordinates in [
        ("EDUCACION", ["B295", "E295", "B296", "E296", "B297"]),
        ("CAPACITACION", ["H295", "L295", "H296", "L296", "H297", "L297"]),
        ("IDIOMA", ["P295", "S295", "P296", "S296", "P297", "S297"]),
    ]:
        for index, coordinate in enumerate(coordinates, start=1):
            name = choice(sheet, coordinate)
            training.append((group, bounded_code(group, name, index, 40), name))
    assert_count("personnel training types", training, 17)

    return {
        "transports": transports,
        "frequencies": frequencies,
        "criteria": criteria,
        "plant": plant,
        "facility_categories": facility_categories,
        "facility_types": facility_types,
        "complementary": complementary,
        "conservation_states": conservation_states,
        "factors": factors,
        "basic_categories": [(code, name) for code, name, _ in basic_categories],
        "basic_types": basic_types,
        "signage": signage,
        "signage_materials": signage_materials,
        "health_services": health_services,
        "security_services": security_services,
        "communication": communication,
        "threats": threats,
        "policy_questions": policy_questions,
        "activities": activities,
        "promotion_media": promotion_media,
        "training": training,
    }


def emit_simple(lines: list[str], table: str, columns: list[str], rows: list[tuple[object, ...]], updates: list[str]) -> None:
    lines.extend([
        f"INSERT INTO {table} ({', '.join(columns)})",
        "VALUES",
        sql_values(rows, len(columns)),
        "ON CONFLICT (codigo) DO UPDATE SET",
    ])
    lines.append(",\n".join(f"  {update}" for update in updates) + ";")
    lines.append("")


def emit_parented(
    lines: list[str],
    table: str,
    source_columns: list[str],
    target_columns: list[str],
    parent_table: str,
    parent_source_column: str,
    parent_target_column: str,
    rows: list[tuple[object, ...]],
    updates: list[str],
) -> None:
    lines.extend([
        f"INSERT INTO {table} ({', '.join(target_columns)})",
        "SELECT parent.id, source." + ", source.".join(source_columns[1:]),
        "FROM (VALUES",
        sql_values(rows, len(source_columns)),
        f") AS source({', '.join(source_columns)})",
        f"JOIN {parent_table} AS parent ON parent.{parent_target_column} = source.{parent_source_column}",
        "ON CONFLICT (codigo) DO UPDATE SET",
    ])
    lines.append(",\n".join(f"  {update}" for update in updates) + ";")
    lines.append("")


def emit_check(lines: list[str], table: str, codes: list[str], expected: int) -> None:
    code_sql = ", ".join(sql_literal(code) for code in codes)
    lines.extend([
        f"IF (SELECT COUNT(*) FROM {table} WHERE activo AND codigo IN ({code_sql})) <> {expected} THEN",
        f"  RAISE EXCEPTION 'XLSM seed check failed for {table}: expected {expected} active source rows';",
        "END IF;",
    ])


def build_sql(source: Path, catalogs: dict[str, list[tuple[object, ...]]]) -> str:
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    lines = [
        "-- Catálogos operativos fijos derivados del XLSM institucional de la ficha turística.",
        f"-- Fuente: {source.as_posix()}",
        f"-- SHA-256: {digest}",
        "-- No carga datos de ejemplo, clima libre, materiales de vía, localidades ni zonas turísticas.",
        "-- No elimina opciones técnicas administradas posteriormente; solo actualiza/reactiva códigos fuente.",
        "",
        "BEGIN;",
        "",
        "SET LOCAL lock_timeout = '5s';",
        "SET LOCAL statement_timeout = '120s';",
        "",
    ]

    emit_simple(
        lines,
        "tipos_accesibilidad",
        ["codigo", "nombre"],
        [
            ("GENERAL", "Accesibilidad general"),
            ("FISICA", "Discapacidad física"),
            ("VISUAL", "Discapacidad visual"),
            ("AUDITIVA", "Discapacidad auditiva"),
            ("COGNITIVA", "Discapacidad intelectual o psicosocial"),
        ],
        ["nombre = EXCLUDED.nombre", "activo = TRUE"],
    )
    emit_simple(lines, "tipos_transporte", ["codigo", "nombre"], catalogs["transports"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "frecuencias_servicio", ["codigo", "nombre"], catalogs["frequencies"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])

    emit_parented(
        lines,
        "criterios_accesibilidad",
        ["tipo_codigo", "codigo", "descripcion", "orden"],
        ["tipo_accesibilidad_id", "codigo", "descripcion", "orden"],
        "tipos_accesibilidad",
        "tipo_codigo",
        "codigo",
        catalogs["criteria"],
        [
            "tipo_accesibilidad_id = EXCLUDED.tipo_accesibilidad_id",
            "descripcion = EXCLUDED.descripcion",
            "orden = EXCLUDED.orden",
            "activo = TRUE",
        ],
    )

    emit_simple(
        lines,
        "tipos_planta_turistica",
        ["grupo", "codigo", "nombre", "unidad_1", "unidad_2", "unidad_3"],
        catalogs["plant"],
        ["grupo = EXCLUDED.grupo", "nombre = EXCLUDED.nombre", "unidad_1 = EXCLUDED.unidad_1", "unidad_2 = EXCLUDED.unidad_2", "unidad_3 = EXCLUDED.unidad_3", "activo = TRUE"],
    )
    emit_simple(lines, "categorias_facilidad", ["codigo", "nombre"], catalogs["facility_categories"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    lines.extend([
        "-- Reconciliación de aliases de cargas previas; se conserva el historial.",
        "UPDATE tipos_facilidad AS legacy",
        "SET activo = FALSE",
        "WHERE legacy.codigo IN (",
        "  'GESTION_PUNTO_DE_INFORMACION',",
        "  'GESTION_I_TUR',",
        "  'GESTION_CENTRO_DE_INTERPRETACION',",
        "  'GESTION_CENTRO_DE_FACILITACION_TURISTICA',",
        "  'GESTION_05',",
        "  'OBSERVACION_GARITAS_DE_GUARDIANIA',",
        "  'OBSERVACION_MIRADORES',",
        "  'OBSERVACION_08',",
        "  'OBSERVACION_09',",
        "  'RECORRIDO_SENDEROS',",
        "  'RECORRIDO_11',",
        "  'RECORRIDO_AREAS_DE_ACAMPAR',",
        "  'RECORRIDO_REFUGIO_DE_ALTA_MONTANA',",
        "  'SERVICIO_BATERIAS_SANITARIAS',",
        "  'SERVICIO_ESTACIONAMIENTOS'",
        ")",
        "AND legacy.activo",
        "AND EXISTS (",
        "  SELECT 1",
        "  FROM tipos_facilidad AS canonical",
        "  WHERE canonical.codigo IN (",
        "    'PUNTO_INFORMACION', 'I_TUR', 'CENTRO_INTERPRETACION',",
        "    'CENTRO_FACILITACION', 'CENTRO_RECEPCION',",
        "    'GARITA', 'MIRADOR', 'TORRE_AVISTAMIENTO', 'TORRE_SALVAVIDAS',",
        "    'SENDERO', 'SOMBRA_DESCANSO', 'AREA_ACAMPAR', 'REFUGIO_MONTAÑA',",
        "    'BATERIA_SANITARIA', 'ESTACIONAMIENTO'",
        "  )",
        ");",
        "",
    ])
    emit_parented(
        lines,
        "tipos_facilidad",
        ["categoria_codigo", "codigo", "nombre"],
        ["categoria_facilidad_id", "codigo", "nombre"],
        "categorias_facilidad",
        "categoria_codigo",
        "codigo",
        catalogs["facility_types"],
        ["categoria_facilidad_id = EXCLUDED.categoria_facilidad_id", "nombre = EXCLUDED.nombre", "activo = TRUE"],
    )
    emit_simple(lines, "tipos_servicio_complementario", ["codigo", "nombre"], catalogs["complementary"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])

    emit_simple(lines, "estados_conservacion", ["codigo", "nombre"], catalogs["conservation_states"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "factores_alteracion", ["origen", "codigo", "nombre"], catalogs["factors"], ["origen = EXCLUDED.origen", "nombre = EXCLUDED.nombre", "activo = TRUE"])

    emit_simple(lines, "categorias_servicio_basico", ["codigo", "nombre"], catalogs["basic_categories"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_parented(
        lines,
        "tipos_servicio_basico",
        ["categoria_codigo", "codigo", "nombre"],
        ["categoria_servicio_basico_id", "codigo", "nombre"],
        "categorias_servicio_basico",
        "categoria_codigo",
        "codigo",
        catalogs["basic_types"],
        ["categoria_servicio_basico_id = EXCLUDED.categoria_servicio_basico_id", "nombre = EXCLUDED.nombre", "activo = TRUE"],
    )
    emit_simple(lines, "tipos_senaletica", ["ambiente", "codigo", "nombre"], catalogs["signage"], ["ambiente = EXCLUDED.ambiente", "nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "materiales_senaletica", ["codigo", "nombre"], catalogs["signage_materials"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "tipos_servicio_salud", ["codigo", "nombre"], catalogs["health_services"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "tipos_servicio_seguridad", ["codigo", "nombre"], catalogs["security_services"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "tipos_comunicacion", ["grupo", "codigo", "nombre"], catalogs["communication"], ["grupo = EXCLUDED.grupo", "nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "tipos_amenaza", ["codigo", "nombre"], catalogs["threats"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "preguntas_politica", ["codigo", "pregunta", "orden"], catalogs["policy_questions"], ["pregunta = EXCLUDED.pregunta", "orden = EXCLUDED.orden", "activo = TRUE"])

    emit_parented(
        lines,
        "grupos_actividad",
        ["categoria_codigo", "codigo", "nombre"],
        ["categoria_atractivo_id", "codigo", "nombre"],
        "categorias_atractivo",
        "categoria_codigo",
        "codigo",
        [
            ("AN", "AGUA", "Agua"),
            ("AN", "AIRE", "Aire"),
            ("AN", "TIERRA", "Tierra"),
            ("MC", "CULTURA", "Actividades culturales"),
        ],
        ["categoria_atractivo_id = EXCLUDED.categoria_atractivo_id", "nombre = EXCLUDED.nombre", "activo = TRUE"],
    )
    emit_parented(
        lines,
        "actividades_turisticas",
        ["grupo_codigo", "codigo", "nombre"],
        ["grupo_actividad_id", "codigo", "nombre"],
        "grupos_actividad",
        "grupo_codigo",
        "codigo",
        catalogs["activities"],
        ["grupo_actividad_id = EXCLUDED.grupo_actividad_id", "nombre = EXCLUDED.nombre", "activo = TRUE"],
    )
    emit_simple(lines, "tipos_medio_promocion", ["codigo", "nombre"], catalogs["promotion_media"], ["nombre = EXCLUDED.nombre", "activo = TRUE"])
    emit_simple(lines, "tipos_formacion_personal", ["grupo", "codigo", "nombre"], catalogs["training"], ["grupo = EXCLUDED.grupo", "nombre = EXCLUDED.nombre", "activo = TRUE"])

    lines.extend(["-- Verificaciones de presencia de todos los códigos fuente.", "DO $$", "BEGIN"])
    checks = [
        ("tipos_transporte", catalogs["transports"], 14),
        ("frecuencias_servicio", catalogs["frequencies"], 4),
        ("criterios_accesibilidad", catalogs["criteria"], 51),
        ("tipos_planta_turistica", catalogs["plant"], 20),
        ("categorias_facilidad", catalogs["facility_categories"], 5),
        ("tipos_facilidad", catalogs["facility_types"], 16),
        ("tipos_servicio_complementario", catalogs["complementary"], 5),
        ("estados_conservacion", catalogs["conservation_states"], 4),
        ("factores_alteracion", catalogs["factors"], 22),
        ("categorias_servicio_basico", catalogs["basic_categories"], 4),
        ("tipos_servicio_basico", catalogs["basic_types"], 23),
        ("tipos_senaletica", catalogs["signage"], 24),
        ("materiales_senaletica", catalogs["signage_materials"], 3),
        ("tipos_servicio_salud", catalogs["health_services"], 5),
        ("tipos_servicio_seguridad", catalogs["security_services"], 4),
        ("tipos_comunicacion", catalogs["communication"], 8),
        ("tipos_amenaza", catalogs["threats"], 8),
        ("preguntas_politica", catalogs["policy_questions"], 4),
        ("actividades_turisticas", catalogs["activities"], 56),
        ("tipos_medio_promocion", catalogs["promotion_media"], 8),
        ("tipos_formacion_personal", catalogs["training"], 17),
    ]
    for table, rows, expected in checks:
        code_index = 1 if table in {
            "criterios_accesibilidad",
            "tipos_planta_turistica",
            "tipos_facilidad",
            "factores_alteracion",
            "tipos_servicio_basico",
            "tipos_senaletica",
            "tipos_comunicacion",
            "tipos_formacion_personal",
            "actividades_turisticas",
        } else 0
        if table == "preguntas_politica":
            code_index = 0
        emit_check(lines, table, [str(row[code_index]) for row in rows], expected)
    lines.extend(["END;", "$$;", "", "COMMIT;", ""])
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    workbook = load_workbook(args.source, data_only=False)
    catalogs = extract_operational_catalogs(workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(build_sql(args.source, catalogs), encoding="utf-8")
    print(f"Wrote {args.output} from {args.source}")


if __name__ == "__main__":
    main()
